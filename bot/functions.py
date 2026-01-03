import asyncio
import json
from datetime import datetime,timedelta
import json
import requests
from requests.auth import HTTPBasicAuth
import pytz
import aiohttp
import json
from datetime import datetime
import asyncio
import ssl
import psutil
import os
import asyncio
import win32print
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import barcode as barcode
from barcode.writer import ImageWriter
import win32com.client
import psutil
import pythoncom



uzbekistan_tz = pytz.timezone('Asia/Tashkent')
today = datetime.now(uzbekistan_tz).strftime('%d.%m.%Y')



def get_info_for_product():
    url = "https://smartup.online/b/anor/mxsx/mr/inventory$export"

    auth = HTTPBasicAuth('ramazon14@falcon', '111222')

    headers = {
        'filial_id': '5012602',
        'project_id': 'trade',
        'Content-Type': 'application/json'
    }

    data = {
        "code": "",
        "begin_created_on": "",
        "end_created_on": "",
        "begin_modified_on": "",
        "end_modified_on": ""
    }

    response = requests.get(url, json=data, auth=auth, headers=headers)

    if response.status_code == 200:
        
        # JSON ma'lumotlarini dekodlash
        response_dict = response.json()

        # Ma'lumotlarni tartibli JSON faylga yozish va Unicode belgilarni to'g'ri ko'rsatish
        with open('data.json', 'w', encoding='utf-8') as file:
            json.dump(response_dict, file, indent=4, ensure_ascii=False)
        return True
    else:
        return False



def get_price_for_product():
    url = "https://smartup.online/b/anor/mxs/mkf/product_price$export"

    auth = HTTPBasicAuth('ramazon14@falcon', '111222')

    headers = {
        'filial_id': '5012602',
        'project_code': 'trade',
    }

    data = {
        
    }

    response = requests.get(url, json=data, auth=auth, headers=headers)

    if response.status_code == 200:
        response_dict = response.json()
        with open('prices.json', 'w', encoding='utf-8') as file:
            json.dump(response_dict, file, indent=4, ensure_ascii=False)
        return True
    else:
        return False
def do_all():
    info = get_info_for_product()
    price = get_price_for_product()
    if price and info:
        with open('data.json', 'r', encoding='utf-8') as data_file:
            data = json.load(data_file)

        with open('prices.json', 'r', encoding='utf-8') as prices_file:
            prices = json.load(prices_file)

        price_dict = {}
        for item in prices["inventory"]:
            price_uzs = item["price_type"][0]["price"] if len(item["price_type"]) > 0 else None
            price_usd = item["price_type"][1]["price"] if len(item["price_type"]) > 1 else None
            price_dict[item["inventory_code"]] = {
                "price_uzs": price_uzs,
                "price_usd": price_usd
            }

        for product in data["inventory"]:
            code = product["code"]  
            if code in price_dict:  
                product["price_uzs"] = price_dict[code]["price_uzs"]
                product["price_usd"] = price_dict[code]["price_usd"]
            else:
                product["price_uzs"] = None
                product["price_usd"] = None

        with open('updated_data.json', 'w', encoding='utf-8') as updated_file:
            json.dump(data, updated_file, ensure_ascii=False, indent=4)
        return True
    else:
        return False


async def get_today_orders():
    url = "https://smartup.online/b/trade/txs/tdeal/order$export"
    
    auth = aiohttp.BasicAuth('ramazon14@falcon', '111222')

    headers = { 
        'filial_id': '5012602',
        'project_code': 'trade',
    }

    data = {
        
        "filial_codes": [{"filial_code": "5012602"}],
        "statuses":[],
      
        "begin_deal_date": today,
        "end_deal_date": today,
        }

    ssl_context = ssl.create_default_context()
    ssl_context.check_hostname = False
    ssl_context.verify_mode = ssl.CERT_NONE  

    async with aiohttp.ClientSession(auth=auth) as session:
        async with session.post(url, json=data, headers=headers, ssl=ssl_context) as response:
            text = await response.text()  # Javobni text sifatida olish
            
            
            if response.status == 200:
                try:
                    json_data = json.loads(text)  # JSON ga o'girishga harakat qilish
                    with open('today_orders.json', 'w', encoding='utf-8') as file:
                        json.dump(json_data, file, indent=4, ensure_ascii=False)
                    update_orders(order_file="today_orders.json", output_file="today_orders.json")
                    
                    return True

                except json.JSONDecodeError:
                    return False
            else:
                return False  
async def get_lastday_orders():
    yesterday = (datetime.now(uzbekistan_tz) - timedelta(days=1)).strftime('%d.%m.%Y')

    url = "https://smartup.online/b/trade/txs/tdeal/order$export"
    auth = aiohttp.BasicAuth('ramazon14@falcon', '111222')

    headers = { 
        'filial_id': '5012602',
        'project_code': 'trade',
    }

    data = {
        "filial_codes": [{"filial_code": "5012602"}],
        "statuses":[],
        "begin_deal_date": yesterday,
        "end_deal_date": yesterday,
    }

    ssl_context = ssl.create_default_context()
    ssl_context.check_hostname = False
    ssl_context.verify_mode = ssl.CERT_NONE  

    async with aiohttp.ClientSession(auth=auth) as session:
        async with session.post(url, json=data, headers=headers, ssl=ssl_context) as response:
            text = await response.text()
            
            if response.status == 200:
                try:
                    json_data = json.loads(text)  # JSON formatga o‘girish
                    with open('lastday_orders.json', 'w', encoding='utf-8') as file:
                        json.dump(json_data, file, indent=4, ensure_ascii=False)

                    update_orders(order_file="lastday_orders.json", output_file="lastday_orders.json")
                    return True
                except json.JSONDecodeError:
                    return False

            else:
                return False
async def get_orders():
    url = "https://smartup.online/b/trade/txs/tdeal/order$export"
    auth = aiohttp.BasicAuth('ramazon14@falcon', '111222')

    headers = { 
        'filial_id': '5012602',
        'project_code': 'trade',
    }

    data = {
        "filial_codes": [{"filial_code": "5012602"}],
        "begin_deal_date": "10.01.2025", 
        "end_deal_date": today,  
    }

    ssl_context = ssl.create_default_context()
    ssl_context.check_hostname = False
    ssl_context.verify_mode = ssl.CERT_NONE  

    async with aiohttp.ClientSession(auth=auth) as session:
        async with session.post(url, json=data, headers=headers, ssl=ssl_context) as response:
            text = await response.text()

            if response.status == 200:
                try:
                    json_data = json.loads(text)  # JSON formatga o‘girish
                    with open('orders.json', 'w', encoding='utf-8') as file:
                        json.dump(json_data, file, indent=4, ensure_ascii=False)
                    update_orders(order_file="orders.json", output_file="orders.json")
                    return True
                except json.JSONDecodeError:
                    return False
            else:
                return False



def update_orders(order_file, output_file):
    # Yagona o'qish va saqlashni optimallashtirish
    def load_json(file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    
    def save_json(data, file_path):
        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)
    
    # data.json ni bir marta o'qiymiz
    data = load_json('data.json')
    
    # product_id va box_quant ni bitta lug'atga saqlash
    product_boxes = {item["product_id"]: int(item["box_quant"]) if item.get("box_quant") else None for item in data["inventory"]}

    # order_file ni o'qiymiz
    orders = load_json(order_file)

    # Har bir orderda mahsulotlar uchun box_quant ni yangilaymiz
    for order in orders["order"]:
        for product in order["order_products"]:
            product_id = product["product_id"]
            product["box_quant"] = product_boxes.get(product_id, None)
    
    # Yangilangan ma'lumotni output_file ga saqlaymiz
    save_json(orders, output_file)



async def get_today_orders_by_manager(sales_manager_id):
    await get_today_orders()
    orders_file = 'today_orders.json'
    def load_json(file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    today_date = datetime.today().strftime("%d.%m.%Y")  
    orders = load_json(orders_file)
    
   
    filtered_orders = []
    for order in orders["order"]:
        if order["sales_manager_id"] == sales_manager_id:
            filtered_orders.append(order)
    
    return filtered_orders[::-1]

async def get_last_day_orders_by_manager(sales_manager_id):
    await get_lastday_orders()
    orders_file = 'lastday_orders.json'
    def load_json(file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    
    yerterday = (datetime.today() - timedelta(days=1)).strftime("%d.%m.%Y")
    orders = load_json(orders_file)
    
    filtered_orders = []
    for order in orders["order"]:
        if order["sales_manager_id"] == sales_manager_id:
            filtered_orders.append(order)
    
    return filtered_orders[::-1]

async def check_user_saler_id(saler_id):
    orders_file = 'orders.json'
    with open(orders_file, "r", encoding="utf-8") as f:
        salers = json.load(f)
    checking = salers['order']
    for saler in checking:
        if saler['sales_manager_id']==str(saler_id):
            return True
    return False

async def today_all_orders():
    orders_file = 'today_orders.json'
    def load_json(file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    
    orders = load_json(orders_file)["order"]
    
    # Buyurtmalarni deal_time bo‘yicha saralash (eng yangilari oldin)
    sorted_orders = sorted(orders, key=lambda x: datetime.strptime(x["deal_time"], "%d.%m.%Y %H:%M:%S"), reverse=True)

    return sorted_orders
async def yerterday_all_orders():
    orders_file = 'lastday_orders.json'

    def load_json(file_path):
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    
    orders = load_json(orders_file)["order"]
    
    sorted_orders = sorted(orders, key=lambda x: datetime.strptime(x["deal_time"], "%d.%m.%Y %H:%M:%S"), reverse=True)


    return sorted_orders


async def get_today_order_info_by_deal_id(deal_id_to_find):
    file_path = 'today_orders.json' 
    with open(file_path, 'r', encoding='utf-8') as f:
        orders_data = json.load(f)
    for order in orders_data['order']:
        if order['deal_id'] == deal_id_to_find:
            return {
                'deal_id': order['deal_id'],
                'deal_time': order['deal_time'],
                'delivery_date': order['delivery_date'],
                'booked_date': order['booked_date'],
                'total_amount': order['total_amount'],
                'room_name': order['room_name'],
                'sales_manager_name': order['sales_manager_name'],
                'sales_manager_id': order['sales_manager_id'],
                'person_name': order['person_name'],
                'currency_code': order['currency_code'],
                'status': order['status'],
                'payment_type_code': order['payment_type_code'],
                'deal_margin_kind': order['deal_margin_kind'],
                'deal_margin_value': order['deal_margin_value'],
                'note': order['note'],
                'deal_note': order['deal_note'],
                'with_marking': order['with_marking'],
                'self_shipment': order['self_shipment'],
                'total_weight_netto': order['total_weight_netto'],
                'total_weight_brutto': order['total_weight_brutto'],
                'total_litre': order['total_litre'],
                'order_products': order['order_products']
            }
    
    return None  # deal_id topilmasa
async def get_lastday_order_info_by_deal_id(deal_id_to_find):
    file_path = 'lastday_orders.json'  #

    # JSON faylini ochish va ma'lumotlarni o'qish
    with open(file_path, 'r', encoding='utf-8') as f:
        orders_data = json.load(f)
    
    # Zakazni topish
    for order in orders_data['order']:
        if order['deal_id'] == deal_id_to_find:
            return {
                'deal_id': order['deal_id'],
                'deal_time': order['deal_time'],
                'delivery_date': order['delivery_date'],
                'booked_date': order['booked_date'],
                'total_amount': order['total_amount'],
                'room_name': order['room_name'],
                'sales_manager_id': order['sales_manager_id'],
                'sales_manager_name': order['sales_manager_name'],
                'person_name': order['person_name'],
                'currency_code': order['currency_code'],
                'status': order['status'],
                'payment_type_code': order['payment_type_code'],
                'deal_margin_kind': order['deal_margin_kind'],
                'deal_margin_value': order['deal_margin_value'],
                'note': order['note'],
                'deal_note': order['deal_note'],
                'with_marking': order['with_marking'],
                'self_shipment': order['self_shipment'],
                'total_weight_netto': order['total_weight_netto'],
                'total_weight_brutto': order['total_weight_brutto'],
                'total_litre': order['total_litre'],
                'order_products': order['order_products']
            }
    
    return None  
async def get_order_info_by_deal_id(deal_id_to_find):
    await get_orders()

    file_path = 'orders.json'
    with open(file_path, 'r', encoding='utf-8') as f:
        orders_data = json.load(f)
    
    for order in orders_data['order']:
        if order['deal_id'] == deal_id_to_find:
            return {
                'deal_id': order['deal_id'],
                'deal_time': order['deal_time'],
                'delivery_date': order['delivery_date'],
                'booked_date': order['booked_date'],
                'total_amount': order['total_amount'],
                'room_name': order['room_name'],
                'sales_manager_id': order['sales_manager_id'],
                'sales_manager_name': order['sales_manager_name'],
                'person_name': order['person_name'],
                'currency_code': order['currency_code'],
                'status': order['status'],
                'payment_type_code': order['payment_type_code'],
                'deal_margin_kind': order['deal_margin_kind'],
                'deal_margin_value': order['deal_margin_value'],
                'note': order['note'],
                'deal_note': order['deal_note'],
                'with_marking': order['with_marking'],
                'self_shipment': order['self_shipment'],
                'total_weight_netto': order['total_weight_netto'],
                'total_weight_brutto': order['total_weight_brutto'],
                'total_litre': order['total_litre'],
                'order_products': order['order_products']
            }
    
    return None  

async def get_selected_orders(deal_id):
    url = "https://smartup.online/b/trade/txs/tdeal/order$export"
    
    auth = aiohttp.BasicAuth('ramazon14@falcon', '111222')

    headers = { 
        'filial_id': '5012602',
        'project_code': 'trade',
    }

    data = {
        "filial_codes": [{"filial_code": "5012602"}],
        "deal_id": deal_id,
        "statuses":[],
        "begin_deal_date": today,
        "end_deal_date": today,
        }

    ssl_context = ssl.create_default_context()
    ssl_context.check_hostname = False
    ssl_context.verify_mode = ssl.CERT_NONE  

    async with aiohttp.ClientSession(auth=auth) as session:
        async with session.post(url, json=data, headers=headers, ssl=ssl_context) as response:
            text = await response.text() 
            
            
            if response.status == 200:
                try:
                    json_data = json.loads(text)  
                    print(json_data)
                    return json_data

                except json.JSONDecodeError:
                    return False
            else:
                return False  

async def update_document(template_path, new_file_path, name, price, usd_price, barcode_image_path):
    """ Hujjatni yangilash va saqlash """
    doc = await asyncio.to_thread(Document, template_path)
    for para in doc.paragraphs:
        if "Mahsulotnomi" in para.text:
            para.text = ""
            run = para.add_run(f"{name} / {usd_price}")
            run.font.size = Pt(7)
            run.bold = True
            para.paragraph_format.left_indent = Cm(0.5)
        if "50000" in para.text:
            para.text = ""
            run = para.add_run(f"   {price} сум")
            run.font.size = Pt(14)
            run.bold = True
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    if os.path.exists(barcode_image_path):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(barcode_image_path, width=Cm(4.2), height=Cm(1.5))
    
    os.makedirs("documents", exist_ok=True)
    save_path = os.path.join("documents", new_file_path)
    await asyncio.to_thread(doc.save, save_path)
    return save_path

async def print_barcode(word_name, data_to_encode, name, price, usd_price, barcode_name='barcode.png', page=1):
    try:
        new_file_path = f'{word_name}.docx'
        print(data_to_encode,len(data_to_encode))
        barcode_image_path = await asyncio.to_thread(generate_barcode, data_to_encode, barcode_name)
        saved_file_path = await update_document('aaa.docx', new_file_path, name, price, usd_price, barcode_image_path)
        await print_document(saved_file_path, page)
        os.remove(f"{barcode_name}.png")
        os.remove(f"documents/{word_name}.docx")
        return True
    except Exception as e:
        print(e)
        if os.path.exists(f"{barcode_name}.png"):
            os.remove(f"{barcode_name}.png")
        return False

async def print_document(file_path, copies=1):
    kill_task()
    file_path = os.path.abspath(file_path)
    
    if os.path.exists(file_path):
        try:
            loop = asyncio.get_event_loop()
            await loop.run_in_executor(None, _print_document_sync, file_path, copies)
        except Exception as e:
            print(f"Chop etishda xatolik yuz berdi: {e}")
    else:
        print("Fayl topilmadi!")  

def _print_document_sync(file_path, copies):
    """ Hujjatni sinxron ravishda chop etish """
    try:
        pythoncom.CoInitialize()

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(file_path)

        word.ActivePrinter = "Xprinter XP-303B"

        for _ in range(copies):
            doc.PrintOut()
        doc.Close(False)
        word.Quit()

        print(f"✅ {file_path} fayli {copies} nusxada Xprinter XP-303B (копия 1) printeridan chop etish uchun yuborildi.")
        pythoncom.CoUninitialize()
        return True
    except Exception as e:
        print(f"❌ Xatolik: {e}")
        return False  



def kill_task():
    task_name = 'WORD.EXE'
    for proc in psutil.process_iter(['pid', 'name']):
        
        if task_name.lower() in proc.info['name'].lower(): 
            try:
                proc.kill()  
            except psutil.NoSuchProcess:
                pass
            except psutil.AccessDenied:
                pass

def generate_barcode(data, filename='barcode.png'):
    """ Shtrixkod yaratish va saqlash """
    barcode_class = barcode.get_barcode_class('ean13')
    barcode_obj = barcode_class(data, writer=ImageWriter())
    return barcode_obj.save(filename, options={
        "module_height": 15.0,
        "font_size": 10,
        "dpi": 650,
        "text_distance": 5.0,
    })


async def get_product_by_id(user_input):
    try:
        with open('updated_data.json', 'r', encoding='utf-8') as file:
            data = json.load(file)
    except FileNotFoundError:
        return "Xatolik: updated_data.json fayli topilmadi!"
    except json.JSONDecodeError:
        return "Xatolik: JSON fayl noto‘g‘ri formatda!"

    found_products = []

    for product in data.get("inventory", []):
        product_id = str(product.get("product_id", ""))
        if user_input == product_id:  
            found_products.append(product)

    if found_products:
        return found_products
    else:
        print("Hech narsa topilmadi!")
        return None